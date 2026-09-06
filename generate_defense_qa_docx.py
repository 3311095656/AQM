from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("AQMV2_复赛答辩问题与参考答案.docx")


def qa(probability, question, oral, simple, boundary):
    return {
        "probability": probability,
        "question": question,
        "oral": oral,
        "simple": simple,
        "boundary": boundary,
    }


SECTIONS = [
    (
        "第一部分  极高概率：必须熟练口述",
        [
            qa(
                "极高",
                "请用30秒介绍你们的项目。",
                "我们的项目名称是AQMV2，基于RT-Thread星火1号开发板，核心主控为STM32F407ZGT6。系统采集室内外温湿度、光照、人体存在和可燃气体告警信息，在设备端完成LCD与OLED显示、窗户和风扇控制以及燃气安全联动；联网后通过OneNET完成属性上报、历史查询和远程控制。项目的核心特点是安全控制留在设备端，即使网络中断，燃气告警和本地联动仍能继续运行。",
                "它不是只把温度传到网页，而是能在本地根据环境自动做动作，并把数据上传到云端。",
                "不要把MQ-5说成浓度检测，也不要把模拟空调说成真实家电。",
            ),
            qa(
                "极高",
                "项目解决的核心问题是什么？",
                "普通环境监测设备往往只负责采集和显示，真正发生异常时仍需要人工处理。我们希望把感知、判断和执行放在同一个端侧系统中：平时根据温湿度、光照和人体存在进行舒适性控制，发生燃气告警时由设备本地立即执行高优先级安全动作，云端则负责远程查看、配置和控制。",
                "普通温度计只告诉你热不热，我们的系统还会根据规则控制设备，并在危险时自动处理。",
                "不要声称已经替代成熟智能家居或认证安全设备。",
            ),
            qa(
                "极高",
                "你们的核心创新是什么？传感器和MQTT都是现成的。",
                "我们的创新主要是系统级协同，而不是发明新传感器。第一，建立了燃气安全高于定时、手动和自动控制的优先级；第二，把人体存在和光照共同用于窗户策略，而不是单一阈值控制；第三，把安全控制保留在STM32端，OneNET只承担远程管理，避免网络成为安全功能的单点依赖。我们采用的是可解释规则控制，不把它包装成AI算法。",
                "零件可以买到，但如何安排谁先执行、冲突时听谁的、断网时还能不能工作，是系统设计的价值。",
                "不要使用“多传感器融合算法”“人工智能”“自学习”等代码无法支撑的词。",
            ),
            qa(
                "极高",
                "为什么选择RT-Thread星火1号和STM32F407？",
                "星火1号对RT-Thread的BSP和板载资源支持比较完整，带有LCD、Wi-Fi扩展能力和丰富外设接口，适合在比赛周期内快速完成多传感器、多执行器和端云通信。STM32F407具备足够的性能、Flash和I2C、SPI、UART、PWM等资源，可以同时运行RT-Thread、网络协议栈、MQTT和显示任务。竞赛阶段优先考虑开发稳定性和功能完整度，量产阶段再根据实际负载缩减主控和板级资源。",
                "选择这块板主要是为了少花时间重新做底层驱动，把精力放在系统功能和可靠性上。",
                "必须说明使用的是现成星火1号平台，不要说成自研STM32F407开发板。",
            ),
            qa(
                "极高",
                "为什么不用ESP32？",
                "ESP32完全可以完成类似功能，并且集成Wi-Fi，在成本和体积方面有优势。但本项目不仅是联网传感器，还需要同时驱动双屏、多路I2C、PWM舵机、风扇、蜂鸣器和雷达。星火1号的RT-Thread支持和外设资源更符合当前竞赛原型需求，也能降低驱动适配风险。若进入量产并更重视成本、尺寸和无线集成度，ESP32或ESP32-S3会是需要重新评估的方案。",
                "不是说ESP32不好，而是当前开发板和软件生态更适合比赛时间内稳定完成所有功能。",
                "不要回答“ESP32实时性差”或“ESP32不稳定”。",
            ),
            qa(
                "极高",
                "为什么使用RT-Thread，而不是裸机轮询？",
                "系统中存在不同周期和重要程度的任务：燃气线程约每200毫秒检测一次，主控制循环负责采集、显示和执行器，网络线程处理Wi-Fi、NTP和周期上报，Paho还负责MQTT收发。RT-Thread让我们可以通过优先级隔离安全任务和网络任务，并使用互斥锁传递完整的传感器快照。裸机也能实现，但状态机会更复杂，模块耦合和后续扩展成本更高。",
                "RTOS相当于给不同工作安排独立岗位和优先级，避免网络卡住时连报警都不能运行。",
                "不要声称使用RTOS后就自动获得绝对实时性和零死锁。",
            ),
            qa(
                "极高",
                "燃气报警的安全优先级如何保证？",
                "MQ-5由独立高优先级线程约每200毫秒读取一次。检测到告警后，蜂鸣器和红灯立即动作，主控制循环强制开窗、开风扇，并关闭空调与湿控状态。MQTT命令入口也有安全仲裁：报警期间，关风扇、关窗、关闭蜂鸣器、开启空调或湿控、切换自动模式都会被拒绝并返回403。整个安全动作在STM32端完成，不依赖网页和OneNET。",
                "危险发生时，本地安全规则拥有最高决定权，网页不能把安全设备关掉。",
                "这是竞赛原型安全逻辑，不等于通过燃气、消防或防爆认证。",
            ),
            qa(
                "极高",
                "MQ-5能不能测出具体ppm浓度？",
                "当前不能。我们读取的是MQ-5模块的DO数字比较器输出，只能判断气体是否超过硬件设定阈值。代码没有采集模拟电压，也没有完成传感器预热曲线、Rs/R0标定、温湿度补偿和气体交叉敏感性校准。因此准确表述是“可燃气体阈值告警”，不是ppm浓度测量。",
                "现在得到的是“超过警戒线没有”，不是“空气里有多少ppm”。",
                "绝对不能宣称MQ-5已经实现定量浓度检测。",
            ),
            qa(
                "极高",
                "如何证明传感器测量精度？",
                "软件层面，AHT10温湿度和AP3216C光照采用三点中值滤波，DHT22限制实际采样间隔并保留最近有效值。但滤波只能抑制偶发噪声，不能代替校准。整机精度应通过与标准温湿度计和照度计同位置、同步、多点采样，统计平均误差、最大误差和重复性来证明。没有完成的实测数据，我们只引用器件手册标称值，不临场编造。",
                "让项目传感器和标准仪器一起测，再比较差多少，才叫精度证明。",
                "器件标称精度不能直接说成整机实测精度。",
            ),
            qa(
                "极高",
                "真实端云数据链路是什么？",
                "设备通过Wi-Fi直接连接OneNET MQTT服务器，使用物模型property/post主题上报属性，订阅property/set接收控制，并通过set_reply返回处理结果。网页不直接保存OneNET Token，而是访问本地Node服务；Node在服务端携带Token，通过OneNET HTTPS API查询设备状态、实时属性和历史数据，或者下发属性设置。最终架构是“固件直连OneNET，网页通过Node代理访问OneNET”。",
                "设备直接把数据送到OneNET；网页先问本地服务器，本地服务器再安全地代它访问OneNET。",
                "不要再描述早期本地Aedes桥接方案，当前版本已经不采用该架构。",
            ),
            qa(
                "极高",
                "断网后系统还能做什么？",
                "断网不会影响本地传感器采集、屏幕显示、按键和燃气安全联动，因为这些功能都运行在STM32端。断网期间网页远程控制和云端历史不可用，遥测目前也不会离线缓存补传。Wi-Fi支持重试与自动重连，MQTT恢复后继续上报；定时功能还依赖有效系统时间。",
                "没有网络时，本地设备照常工作，只是手机网页看不到，也不能远程控制。",
                "不要声称当前已经实现离线数据补传。",
            ),
            qa(
                "极高",
                "哪些执行器是真实的，哪些只是模拟？",
                "当前真实执行器包括风扇、SG90窗户舵机、蜂鸣器和状态灯，LCD与OLED也是真实显示设备。空调、加湿和除湿没有直接连接大功率家电，而是通过WS2812不同颜色模拟运行状态，用来验证控制策略和远程接口。另外，舵机没有位置反馈、风扇没有转速反馈，因此上报的是软件指令状态，不是经过传感器确认的物理状态。",
                "风扇、舵机和蜂鸣器会真的动作；空调和湿控目前只用灯表示控制结果。",
                "不能说已经控制真实空调或完成执行器闭环。",
            ),
            qa(
                "极高",
                "现场如何证明数据不是预设动画或模拟数据？",
                "现场会按完整链路演示：改变真实光照和人体存在状态，观察LCD、OLED、串口、OneNET和网页同步变化；再从网页控制真实风扇和舵机；使用安全模拟输入触发燃气告警，展示蜂鸣器、红灯、安全联动和危险远程命令被拒绝；最后断开并恢复热点，证明本地功能不受影响且设备能够重新上线。演示时不运行模拟数据脚本。",
                "同时看实物动作、串口日志和网页数据，三者一起变化，就能证明不是动画。",
                "不要在会场使用危险可燃气体进行演示。",
            ),
        ],
    ),
    (
        "第二部分  高概率：技术原理与实现",
        [
            qa(
                "高",
                "系统使用了哪些传感器和接口？",
                "AHT10通过I2C采集室外温湿度，DHT22通过单总线采集室内温湿度，AP3216C通过I2C采集光照，LD2410C通过数字OUT脚提供人体存在状态，MQ-5通过数字DO脚提供燃气阈值告警。接口类型覆盖I2C、单总线和GPIO，执行侧还使用PWM和普通GPIO。",
                "不同传感器用不同通信方式连接到STM32，最后统一整理成一份环境数据。",
                "LD2410C当前没有使用UART高级数据，MQ-5当前没有使用ADC。",
            ),
            qa(
                "高",
                "为什么同时采集室内和室外温湿度？",
                "室内温湿度用于舒适性控制和目标值比较，室外数据用于环境对比、趋势展示和后续策略扩展。这样可以避免把室外变化直接误认为室内控制效果，也为以后增加通风决策提供基础。",
                "室内数据决定房间怎么调，室外数据告诉我们外部环境是什么情况。",
                "当前空调控制已经修正为使用室内温度。",
            ),
            qa(
                "高",
                "为什么使用三点中值滤波？",
                "连续采集三个数并排序，取中间值，可以过滤一次明显偏大或偏小的脉冲异常。它计算量小，适合STM32实时使用。它的作用是提高数据稳定性，而不是修正传感器长期偏差。",
                "三个数里去掉一个最离谱的，保留中间那个。",
                "不要把中值滤波称为多传感器融合或精度校准。",
            ),
            qa(
                "高",
                "DHT22为什么不是每秒真正采样？",
                "DHT22对采样间隔有要求。驱动会限制真正的物理读取频率，间隔不足时返回最近缓存值；读取失败时也会重试并保留最近结果。这避免了过于频繁读取导致通信失败，但后续还应增加数据年龄和有效性标志。",
                "主界面每秒刷新，不代表传感器每秒都重新测量；有时显示的是最近一次有效数据。",
                "不能把缓存值说成新的物理采样。",
            ),
            qa(
                "高",
                "普通环境传感器初始化失败怎么办？",
                "MQ-5安全线程在普通传感器之前启动，因此环境传感器故障不会让安全告警停止。初始化失败时，系统暂停依赖温湿度和光照的自动控制，并每30秒尝试重新初始化。当前故障时的零值展示仍应被解释为无效状态，而不是实际测量。",
                "普通传感器坏了，系统先保住报警功能，并定期尝试恢复，不会直接死机退出。",
                "不要把故障零值解释为真实0摄氏度或0%湿度。",
            ),
            qa(
                "高",
                "LD2410C读取了哪些信息？",
                "当前只使用LD2410C的OUT引脚获取有人或无人二值状态，没有通过UART读取目标距离、运动能量、静止能量等高级数据。这样实现简单、隐私友好，但无法提供人数和精确位置。",
                "目前雷达只回答“有没有人”，不会告诉我们人在几米远。",
                "不要宣称已实现距离、人数或轨迹检测。",
            ),
            qa(
                "高",
                "窗户自动控制策略是什么？",
                "在自动模式且检测到有人时，系统根据光照把舵机角度分为0、9、18、27、36和45度六档；光照越强，开度越小，用于模拟遮阳和强光抑制；无人时关闭。燃气告警时该舒适性策略被覆盖，窗户强制打开。阈值是原型演示参数，后续应根据实际场景标定。",
                "有人时看光线强弱调角度，无人时关窗，危险时无条件开窗。",
                "不要把演示阈值说成经过建筑舒适度标准验证。",
            ),
            qa(
                "高",
                "温湿度控制策略是什么？为什么不用PID？",
                "温度根据目标值和约±0.5摄氏度中性区选择制冷、制热或待机，湿度根据约±3%中性区选择加湿、除湿或待机。当前执行器主要是离散状态，环境变化缓慢，因此规则控制更透明、容易调试。PID更适合连续可调且有明确动态模型的执行器，当前阶段没有必要为了算法复杂度而使用PID。",
                "温度偏高就制冷、偏低就制热，接近目标就停止；留一段缓冲区避免来回切换。",
                "不要称为PID、模糊控制、自适应或AI。",
            ),
            qa(
                "高",
                "系统如何处理手动、自动、定时和安全控制冲突？",
                "总体原则是燃气安全最高，其次是定时计划和用户控制，最后是自动舒适性控制。远程或按键手动控制会使相应设备退出自动策略；恢复自动后再由规则控制。进入定时计划时保存常用原状态，离开或禁用计划后恢复。燃气告警则无条件覆盖其他模式。",
                "多个控制命令打架时，先听安全规则，再听计划或用户，最后才是普通自动控制。",
                "少见的多设备混合手动状态仍是后续可优化边界。",
            ),
            qa(
                "高",
                "定时计划支持哪些功能？",
                "系统支持最多四个时间段，每段可以独立启用，支持系统自动、温控和湿控三类模式，也支持例如22点到次日6点的跨午夜时段。计划配置由网页通过OneNET下发，设备端按NTP同步后的系统时间判断。",
                "可以设置最多四组开始和结束时间，跨过零点也能判断。",
                "当前计划保存在RAM，设备复位后需要重新配置。",
            ),
            qa(
                "高",
                "进入和退出计划时如何恢复原状态？",
                "进入计划时保存空调和湿控的模式、手动标志以及系统整体自动状态。离开计划或计划被禁用时恢复这些状态。该实现覆盖常见的全自动或全手动场景；如果多个设备分别处于不同的手动和自动组合，单一整体状态无法完全重建，后续应改为保存每个设备的独立模式快照。",
                "定时任务开始前先记住原来状态，结束后恢复；但非常复杂的混合状态还不能百分之百还原。",
                "不要说任何组合都已经实现无损恢复。",
            ),
            qa(
                "高",
                "为什么MQTT使用QoS1？",
                "QoS1要求Broker对消息进行确认，比QoS0的尽力发送更可靠，适合属性上报和控制回复。但QoS1语义是至少一次，网络异常时可能重复，因此量产控制还需要命令ID、幂等判断和最终状态确认。",
                "消息服务器会回复“我收到了”，但同一条消息有可能收到两次。",
                "QoS1不等于执行器一定动作成功。",
            ),
            qa(
                "高",
                "设备如何确认远程命令处理结果？",
                "设备解析OneNET属性设置命令，执行成功后通过set_reply返回成功；非法JSON、缺失参数、错误类型和超长载荷会返回400或413，不安全操作返回403。网页还会检查OneNET业务返回码，并通过后续属性查询观察设备状态。完整产品还应把命令ID、执行状态和超时展示给用户。",
                "不仅要把命令发出去，还要告诉云端是成功、参数错误还是因为安全原因被拒绝。",
                "当前仍不能把软件状态回传等同于物理执行反馈。",
            ),
            qa(
                "高",
                "MQTT消息如何防止缓冲区越界和错误JSON？",
                "回调不再直接修改MQTT底层接收内存，而是先检查长度，再复制到独立缓冲区并补字符串结束符。控制载荷超过缓冲区返回413，JSON无法解析返回400，params还会检查必须为对象，模式和目标值检查必须为数字。这样避免越界写和错误数据被误执行。",
                "先把收到的内容安全地复制到自己的盒子里，确认大小和格式正确后再使用。",
                "固定缓冲区仍有长度上限，不能声称支持任意大的JSON。",
            ),
            qa(
                "高",
                "为什么Node只做代理，不直接连接MQTT？",
                "早期方案让Node和固件使用同一设备身份连接OneNET，会产生Client ID冲突并互相踢下线。当前固件是唯一MQTT设备客户端，Node只负责把网页的HTTPS请求转发给OneNET API，并在服务端保存Token。这样职责清晰，也避免浏览器直接暴露凭据。",
                "设备负责MQTT，服务器负责网页API，两边不再抢同一个账号登录。",
                "答辩必须以当前REST代理架构为准。",
            ),
            qa(
                "高",
                "为什么选择OneNET？",
                "OneNET提供标准物模型、属性上报、属性下发、设备状态和历史查询，能够减少自建云平台的工作量，适合竞赛周期内完成可验证的端云闭环。我们使用平台能力完成接入，但设备端控制、安全逻辑和网页代理均由项目实现。",
                "OneNET提供云端基础设施，我们不用从零开发消息服务器和设备管理平台。",
                "不要把OneNET平台自带能力全部说成团队原创。",
            ),
            qa(
                "高",
                "网页如何保护OneNET Token？",
                "浏览器页面只访问本地/api接口，不再包含OneNET Token。Node服务在服务端添加授权信息，并且设备详情接口只返回status、name和last_time，过滤密钥等不需要展示的字段。竞赛LAN环境暂不实现复杂账号体系，生产环境仍需要登录、权限、TLS和限流。",
                "网页不知道真正密码，只有本地服务器知道，并且服务器只把必要信息交给网页。",
                "服务端源码中的静态Token仍不是量产级安全存储。",
            ),
            qa(
                "高",
                "如何判断设备真的执行了网页命令？",
                "网页首先检查OneNET接口返回的业务码，随后重新查询设备属性，观察固件上报的状态是否变化。对于风扇和舵机，现场还可以直接观察物理动作。不过当前设备状态主要来自软件变量，没有编码器、测速和电流反馈，因此只能证明命令状态和可见动作，不能证明全部执行器已形成闭环。",
                "网页发完命令后再次读取状态，再配合眼睛看到实物动作进行确认。",
                "不要把属性回传说成独立硬件传感器反馈。",
            ),
        ],
    ),
    (
        "第三部分  中概率：工程能力与扩展",
        [
            qa(
                "中",
                "当前固件资源占用如何？",
                "当前构建生成的rtthread.bin约423,016字节，主控内部Flash为1MiB，因此Flash仍有余量。bin大小不能代表全部RAM占用，RAM需要结合map文件、动态堆和RT-Thread线程栈水位分析。答辩时可以报告已确认的Flash数据，但RAM必须以工具实测结果为准。",
                "程序文件大约占芯片Flash的四成，但运行时内存还要单独测。",
                "不要用bin大小冒充RAM占用。",
            ),
            qa(
                "中",
                "板载8MB NOR Flash为什么没有使用？",
                "当前复赛版本优先完成实时采集、安全控制和端云闭环，历史数据暂时由OneNET保存，因此没有为了使用硬件而强行加入Flash逻辑。星火1号是原型验证平台，板载资源不要求全部启用。8MB NOR Flash是后续配置持久化、离线缓存、日志和OTA升级的扩展资源。",
                "开发板自带一块大存储，但当前核心功能不依赖它，所以先保留给后续版本。",
                "必须明确当前版本尚未启用NOR Flash。",
            ),
            qa(
                "中",
                "以后如何使用8MB NOR Flash？",
                "计划通过RT-Thread FAL划分配置区、数据区、日志区和升级下载区。配置区保存计划和校准参数；数据区以环形结构缓存断网遥测；日志区记录复位和故障；升级区保存OTA镜像并配合CRC、签名和回滚。写入时采用RAM批量缓存、按页写入和磨损均衡，避免每两秒擦写同一扇区。",
                "可以把Flash分成几个仓库，分别放设置、断网数据、故障日志和升级文件。",
                "普通NOR Flash不等于安全芯片，不能直接宣称安全保存密钥。",
            ),
            qa(
                "中",
                "软件看门狗能覆盖哪些故障？",
                "当前8秒软件定时器可以处理部分业务流程长时间不喂狗的情况，但主线程和网络线程共享喂狗，任一线程存活可能掩盖另一线程卡死；调度器或中断系统卡死时软件定时器也可能无法执行。工程化方案是各关键线程维护独立心跳，由监督线程确认全部正常后喂硬件IWDG。",
                "现在的看门狗能发现部分卡顿，但不能保证每一种死机都能重启。",
                "不要回答“系统任何死机都会在8秒内复位”。",
            ),
            qa(
                "中",
                "历史数据是否完整？",
                "网页当前展示OneNET单次API请求返回的有限样本，用于竞赛趋势查看。它不等同于完整日统计。若要生成严格的日、周、月最大值和平均值，应增加分页查询，或者由服务端持续拉取后写入时序数据库再聚合。",
                "现在图表只展示取回来的部分数据，不代表一天里每一条数据都参与了统计。",
                "不要把有限样本统计称为完整三日统计。",
            ),
            qa(
                "中",
                "断网期间的数据为什么没有补传？",
                "当前版本没有启用外部Flash和离线消息队列，网络断开时仍执行本地控制，但遥测不会持久化。要实现补传，需要为每条记录增加时间戳和序号，写入环形Flash队列，恢复后限速发送，并处理重复和过期数据。",
                "设备断网时还会工作，但这段时间的云端曲线会有空缺。",
                "不要声称MQTT自动重连会自动补回历史数据。",
            ),
            qa(
                "中",
                "如何扩展到多台设备？",
                "当前产品ID和设备名针对单设备演示。多设备版本需要设备注册表、独立凭据、按设备区分状态和历史、权限控制以及服务端缓存，网页也需要设备选择和分组。Node不应让每个浏览器直接重复请求OneNET，而应统一订阅或缓存后再推送。",
                "现在系统只管理一台设备，多台设备需要给每台建立独立档案和数据空间。",
                "不要称当前代码已经支持大规模设备接入。",
            ),
            qa(
                "中",
                "是否支持OTA升级？",
                "当前未实现OTA。后续可以利用板载NOR Flash保存下载镜像，Bootloader负责版本检查、完整性校验、签名验证、写入内部Flash和失败回滚。OTA必须同时考虑断电保护和密钥安全，不能只实现文件下载。",
                "未来可以从云端下载新固件，但还需要一个可靠的启动程序保证升级失败也能恢复。",
                "OTA属于规划功能，不能放进当前已实现特性。",
            ),
            qa(
                "中",
                "项目功耗如何？",
                "当前系统使用LCD、OLED、Wi-Fi和毫米波雷达持续运行，定位为有线供电环境终端，不是低功耗电池设备。仓库没有完整功耗报告，应该用电流表分别测待机、联网、屏幕点亮和执行器动作状态。若做电池版本，需要关闭屏幕、降低上报频率并使用休眠和唤醒。",
                "这版追求功能完整，不是追求一颗电池用几年。",
                "没有实测数据时不要报具体功耗或续航。",
            ),
            qa(
                "中",
                "项目成本是多少？",
                "当前使用星火1号开发板，属于原型成本而不是量产BOM。答辩时应按真实采购记录列出开发板、传感器、显示器和执行器价格，并区分原型成本与量产预估。量产会根据实际接口减少显示器、选择集成无线主控或设计专用PCB。",
                "开发板版本方便开发但不一定最便宜，真正产品会重新选器件和画板。",
                "不要临场猜测一个没有BOM支持的成本数字。",
            ),
            qa(
                "中",
                "系统安全性是否达到量产要求？",
                "没有。当前版本面向受控竞赛网络，固件和Node采用静态凭据，MQTT仍需要进一步加强TLS和安全存储。量产方案需要设备唯一密钥、Token轮换、TLS、用户登录、角色权限、命令防重放、操作审计和安全升级。",
                "比赛环境里可以工作，但公开部署还需要完整的账号、加密和权限系统。",
                "不要使用“银行级安全”“工业级安全”等无依据表述。",
            ),
            qa(
                "中",
                "为什么Node代理没有用户登录？",
                "当前Node运行在隔离、可信的局域网演示环境，主要目标是隐藏浏览器Token和统一OneNET API，因此没有增加账号系统。生产部署时必须增加登录、会话、角色权限、参数白名单、Origin校验和限流，防止局域网任意用户控制设备。",
                "比赛现场只有受信任的人访问，所以先简化；真正上线必须加门锁和权限。",
                "“LAN可接受”只适用于受控演示，不是永久安全结论。",
            ),
            qa(
                "中",
                "系统如何处理MQTT重复消息？",
                "QoS1可能导致同一消息至少送达一次，也可能重复。当前开关和模式设置大多具有幂等性，重复执行结果相同；但严格产品化仍应记录命令ID、最近处理序号和执行结果，避免非幂等动作重复触发。",
                "同一条开灯命令执行两次通常没影响，但加一次计数或执行一次动作就可能出问题。",
                "不要把QoS1说成“恰好一次”。",
            ),
            qa(
                "中",
                "共享状态是否可能产生竞态？",
                "传感器快照已经使用互斥锁，但部分模式标志和计划配置仍可能被MQTT回调与主线程同时访问。Cortex-M上的单个整数读写通常是原子的，但多个字段组合不一定保持一致。工程化应把远程命令送入消息队列，由统一控制线程修改状态，并为计划配置加锁或双缓冲。",
                "两个线程同时改同一组设置时，单个数可能没坏，但整组状态可能来自不同时间。",
                "不要声称整个系统所有共享状态都已经完全线程安全。",
            ),
            qa(
                "中",
                "如何验证长期稳定性？",
                "应进行24小时或72小时连续运行，记录上报次数、重连次数、堆内存、线程栈水位、传感器失败次数和复位原因；同时进行反复上电、网络抖动和命令压力测试。答辩只报告真实记录，没有测试表时不能回答“长期零故障”。",
                "让设备连续跑几天，并持续记录是否掉线、内存是否越来越少、有没有重启。",
                "稳定性必须由时间和记录证明，不能靠一次演示证明。",
            ),
        ],
    ),
    (
        "第四部分  尖锐追问：边界与工程判断",
        [
            qa(
                "中高",
                "项目是不是功能很多，但技术深度不足？",
                "我们不把模块数量当作技术深度。项目真正的工作集中在RTOS任务划分、跨线程数据一致性、燃气安全仲裁、传感器故障降级、MQTT安全解析、端云架构调整和状态恢复。我们也明确承认没有实现AI、完整AQI和认证安全，这说明项目边界是可验证的。",
                "技术深度不只是用了多少元件，更包括故障时系统会怎么做、冲突时谁优先。",
                "不要用功能数量回避裁判对可靠性和测试证据的追问。",
            ),
            qa(
                "中高",
                "STM32F407和星火1号是不是性能过剩？",
                "对于最终量产产品可能存在优化空间，但竞赛原型需要LCD、Wi-Fi、多个总线、RTOS和快速调试，星火1号能显著降低底层适配风险。当前固件约占1MiB Flash的四成，也保留了功能扩展空间。量产阶段会根据CPU负载、接口数量、功耗和BOM重新选型。",
                "比赛先用资源充足的平台保证完成，产品化再换成更便宜、刚好够用的芯片。",
                "不要回答“资源多所以一定更好”。",
            ),
            qa(
                "中高",
                "MQ-5容易误报，为什么还把它作为安全功能？",
                "当前MQ-5用于验证安全联动机制，而不是作为认证检测器。我们明确其预热、交叉敏感和数字阈值限制。真正产品需要经过标定和认证的传感器、故障诊断以及多条件确认。本项目的价值是证明一旦安全事件成立，端侧优先级和联动路径能够正确执行。",
                "传感器本身不是最终产品级，但它可以用来验证报警后系统是否按正确顺序动作。",
                "不要说MQ-5可以准确识别天然气种类和浓度。",
            ),
            qa(
                "中高",
                "燃气报警时打开普通风扇会不会产生火花？",
                "这是必须承认的工程边界。竞赛原型使用低压风扇验证控制逻辑，不代表实际燃气环境可以直接使用普通风扇。真实部署必须采用符合防爆要求的通风设备、隔离驱动、电源设计和认证方案，并遵循相关安全规范。",
                "逻辑上需要通风，但现实产品的风扇和电路必须是不会引燃气体的安全型号。",
                "不能把竞赛低压风扇方案直接描述为可用于真实燃气泄漏现场。",
            ),
            qa(
                "中高",
                "为什么没有执行器反馈，还称为闭环？",
                "更准确地说，目前是环境状态参与控制的规则闭环，但风扇和舵机自身仍是开环执行。系统知道自己下达了什么命令，却没有独立传感器证明执行器达到目标。产品化应增加舵机位置或限位、风扇转速和电流反馈，并分别上报命令状态和实测状态。",
                "系统能根据环境变化再次调整，但不知道风扇是否真的转、窗户是否被卡住。",
                "答辩时不要笼统说所有设备都已经形成完整闭环。",
            ),
            qa(
                "中高",
                "设备数据为0时如何区分真实0和传感器故障？",
                "当前初始化失败时会暂停自动控制并定期重试，避免零值触发舒适性动作。但显示和上报层仍应进一步加入valid标志、故障码和最后成功时间，让网页显示“--”而不是0。现阶段答辩要把故障零值明确解释为无效占位，而非真实环境值。",
                "安全控制已经不会把故障0值拿去调空调，但页面还应该更明确地写“传感器故障”。",
                "不能把故障占位值当作测量结果统计。",
            ),
            qa(
                "中高",
                "为什么计划配置没有保存到NOR Flash？",
                "当前复赛优先级是完成实时控制和端云闭环，计划由网页快速重新下发，因此暂时放在RAM。我们已经明确该限制，下一阶段会使用FAL和EasyFlash保存计划、阈值和校准参数，并通过版本号、CRC和双备份防止掉电写入损坏。",
                "现在断电后设置会丢，未来会把它写到板载Flash并加校验。",
                "不能把板载Flash预留说成已经实现持久化。",
            ),
            qa(
                "中高",
                "网页显示控制成功，是否代表设备已经执行？",
                "不完全代表。网页首先得到OneNET API受理结果，再通过设备属性回传确认软件状态；真正的物理执行还需要实物观察或执行器反馈。当前风扇和舵机可现场观察，但没有独立反馈传感器。因此我们区分“命令已受理”“设备软件状态已更新”和“物理执行已确认”三个层次。",
                "服务器说收到命令、单片机说状态改变、设备真的动了，是三个不同层次。",
                "不要只凭按钮变色就说控制闭环成功。",
            ),
            qa(
                "中高",
                "你们做过哪些测试？",
                "代码层面已经覆盖正常采集、错误JSON、超长消息、错误类型、不安全命令拒绝、传感器初始化失败和网络恢复路径。现场还应提供真实设备的上报、远控、燃气冲突、断网恢复和连续运行记录。任何精度、响应时间、丢包率和稳定时长，只使用真实测试表中的数字回答。",
                "既要测试正常功能，也要故意断网、传错参数和模拟故障，看系统会不会安全失败。",
                "没有原始记录的指标不要口头报数。",
            ),
            qa(
                "中高",
                "最困难的问题是什么？",
                "最困难的不是读取某一个传感器，而是让本地自动、远程手动、定时计划和燃气安全同时存在又不互相破坏。我们通过明确优先级、在MQTT入口拒绝危险命令、传感器故障时暂停舒适性控制、计划进入和退出时保存恢复状态，逐步把控制逻辑收敛到可解释行为。",
                "难点是很多控制来源同时发命令时，系统仍要知道该听谁的。",
                "回答时应结合自己实际参与的代码，避免只讲团队整体。",
            ),
            qa(
                "中",
                "使用了开源组件，哪些工作是你们自己的？",
                "RT-Thread、Paho MQTT、cJSON、Express和OneNET是成熟基础组件。我们完成了传感器与执行器接入、应用层任务划分、数据结构、控制策略、安全仲裁、OneNET物模型适配、Node代理和Web交互。使用成熟库可以减少重复造轮子，但所有项目级组合、故障处理和验证仍需要自行完成。",
                "操作系统和网络库是工具，如何把它们组合成这个系统以及实现业务逻辑，是团队工作。",
                "不要把开源库代码或OneNET平台功能说成团队原创。",
            ),
            qa(
                "中",
                "如果进入下一轮，最优先改什么？",
                "第一优先是完成标准仪器校准、响应时间、断网恢复和长稳测试，形成量化证据；第二是使用板载NOR Flash实现计划持久化和离线缓存；第三是增加PM2.5、CO2、TVOC和执行器反馈；第四是硬件IWDG、TLS、设备唯一凭据及OTA。升级顺序先补可靠性证据，再扩展功能。",
                "先证明现有功能稳定可信，再增加更多传感器和高级功能。",
                "不要把所有规划功能说成短期内已经可以完成。",
            ),
            qa(
                "中",
                "为什么项目值得晋级？",
                "项目已经形成从感知、端侧控制、安全仲裁到OneNET和Web展示的完整原型，并针对早期的温控变量、传感器故障、MQTT越界、命令安全、Token暴露和页面链路等问题完成了系统性修复。更重要的是，我们能清楚说明已实现能力、原型边界和下一阶段工程路线，项目具备继续深化的基础。",
                "系统不只功能能跑，还经历了问题发现、修复和边界梳理，具备继续做深的条件。",
                "不要用“功能最多”作为唯一晋级理由。",
            ),
            qa(
                "中",
                "团队如何分工，你本人负责什么？",
                "团队可以按硬件接线与驱动、固件架构与控制、OneNET与Web、测试与材料进行说明。我本人负责的内容必须按实际情况回答，并准备解释对应模块的数据流、关键函数、遇到的问题和验证方法。裁判可能直接追问代码，因此不能只背团队统一答案。",
                "说清楚谁做了什么，并且自己负责的部分必须能讲到代码和测试细节。",
                "请在答辩前把本题替换成团队真实分工，绝对不能虚构。",
            ),
        ],
    ),
]


FACTS = [
    ("开发平台", "RT-Thread星火1号开发板"),
    ("核心主控", "STM32F407ZGT6（Cortex-M4）"),
    ("操作系统", "RT-Thread"),
    ("环境传感器", "AHT10、DHT22、AP3216C、LD2410C、MQ-5"),
    ("真实执行器", "风扇、SG90舵机窗、蜂鸣器、状态灯"),
    ("模拟执行器", "空调、加湿、除湿由WS2812状态灯模拟"),
    ("云平台", "OneNET物模型"),
    ("当前架构", "固件直连OneNET MQTT；网页通过Node代理OneNET HTTPS API"),
    ("上报周期", "约2秒一次"),
    ("燃气检测周期", "约200毫秒一次"),
    ("当前固件", "rtthread.bin约423,016字节"),
    ("板载预留资源", "8MB NOR Flash，当前版本尚未启用"),
]


FIXES = [
    ("空调误用室外温度", "改为使用室内温度进行自动控制", "已修复"),
    ("环境传感器失败导致系统退出", "MQ-5优先启动，环境传感器失败时保留安全功能", "已修复"),
    ("传感器故障产生野值", "结构体清零、暂停相关自动控制、每30秒重试", "基本修复"),
    ("计划时段enabled无效", "检查时跳过禁用时段", "已修复"),
    ("计划退出后状态错误", "进入前保存状态，退出或禁用后恢复", "常用场景已修复"),
    ("远程命令覆盖燃气安全", "报警期间拒绝危险命令并返回403", "已修复"),
    ("窗户嵌套布尔值解析错误", "使用统一布尔值解析函数", "已修复"),
    ("MQTT payload越界写", "复制到独立缓冲区并检查长度", "已修复"),
    ("非法JSON和错误类型被接受", "增加400、413及对象和数字类型校验", "已修复"),
    ("MQTT使用QoS0", "属性上报和回复改为QoS1", "已修复"),
    ("Node与固件MQTT身份冲突", "移除Node MQTT设备连接，改为HTTPS API代理", "已修复"),
    ("浏览器暴露OneNET Token", "Token迁移至Node服务端", "已修复"),
    ("设备详情返回敏感字段", "只返回status、name、last_time", "已修复"),
    ("服务器托管旧页面", "统一新页面到server/public", "已修复"),
    ("网页不检查控制结果", "检查OneNET业务码并显示错误", "已修复"),
    ("蜂鸣器跨模块引脚导致编译失败", "移除私有宏跨文件访问并重新构建", "已修复"),
    ("固件构建产物落后于源码", "重新构建最新rtthread.bin", "已修复"),
    ("软件看门狗覆盖不完整", "规划硬件IWDG和各线程心跳", "已知边界"),
    ("计划、日志和离线数据未持久化", "规划使用8MB NOR Flash、FAL和EasyFlash", "后续规划"),
    ("执行器缺少位置/转速反馈", "规划增加限位、编码器、电流或测速", "后续规划"),
]


REDLINES = [
    "MQ-5当前只能进行数字阈值告警，不能输出可靠ppm浓度。",
    "规则控制不是PID、AI、自学习或模糊控制。",
    "空调、加湿和除湿当前由WS2812模拟，不是真实大功率负载。",
    "QoS1只确认Broker收到消息，不等于物理执行器已经完成动作。",
    "当前项目不是经过认证的燃气报警器、消防设备或防爆产品。",
    "板载8MB NOR Flash尚未启用，离线缓存、配置持久化和OTA属于规划功能。",
    "没有测试原始记录时，不报告精度、响应时间、丢包率、功耗和稳定时长的具体数字。",
    "星火1号是现成开发平台，不能描述为团队自研STM32F407开发板。",
]


DEMO_CHECKS = [
    "烧录最新固件，确认构建时间晚于全部源码修改。",
    "冷启动后确认LCD、OLED、传感器和串口日志正常。",
    "按KEY1连接Wi-Fi和OneNET，确认网页设备状态变为在线。",
    "改变光照与人体存在状态，核对实物、串口、OneNET和网页同步变化。",
    "从网页控制真实风扇、舵机窗和蜂鸣器，观察属性回传。",
    "使用安全模拟输入触发燃气告警，验证开窗、开风扇、蜂鸣器和红灯。",
    "报警期间从网页尝试危险命令，确认设备拒绝并返回403。",
    "断开热点，验证本地采集和安全控制继续运行；恢复热点后确认重新上线。",
    "准备串口日志截图、OneNET页面截图和备用演示录像。",
    "现场不启动任何模拟数据发布脚本。",
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_box(doc, title, text, fill, title_color):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(text)
    set_run_font(r2, size=10.5, color=(38, 50, 56))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开Word后右键更新目录"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in (
        ("Title", 28, (17, 63, 103)),
        ("Heading 1", 18, (17, 63, 103)),
        ("Heading 2", 13, (28, 84, 122)),
        ("Heading 3", 11, (55, 71, 79)),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("AQMV2 复赛答辩题库  |  星火1号 · STM32F407 · RT-Thread")
    set_run_font(r, size=8.5, color=(96, 125, 139))
    add_page_number(section.footer.paragraphs[0])

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run("AQMV2")
    set_run_font(r, size=34, bold=True, color=(17, 63, 103))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("嵌入式竞赛复赛答辩问题与参考答案")
    set_run_font(r, size=24, bold=True, color=(28, 84, 122))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("按提问概率排序 · 详细口头回答 · 小白解释 · 答辩边界")
    set_run_font(r, size=12, color=(84, 110, 122))

    cover = doc.add_table(rows=4, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.style = "Table Grid"
    for row, (left, right) in zip(
        cover.rows,
        [
            ("开发平台", "RT-Thread星火1号"),
            ("核心主控", "STM32F407ZGT6"),
            ("文档版本", "复赛候选版 v1.0.0-rc1"),
            ("整理日期", "2026-07-19"),
        ],
    ):
        row.cells[0].text = left
        row.cells[1].text = right
        set_cell_shading(row.cells[0], "DCEAF3")
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10.5, bold=cell is row.cells[0])

    doc.add_paragraph()
    add_box(
        doc,
        "使用说明",
        "先熟练掌握第一部分，再根据本人负责模块学习后续问题。答案必须建立在真实代码和测试记录上；涉及团队分工、成本、精度、功耗、响应时间和稳定性的内容，请在答辩前替换为真实数据。",
        "E8F3F8",
        (17, 94, 130),
    )
    doc.add_page_break()

    doc.add_heading("目录", level=1)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("项目事实速查表", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    hdr.cells[0].text = "项目项"
    hdr.cells[1].text = "当前准确口径"
    set_repeat_table_header(hdr)
    for cell in hdr.cells:
        set_cell_shading(cell, "1C547A")
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=10, bold=True, color=(255, 255, 255))
    for key, value in FACTS:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            set_cell_margins(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9.5, bold=cell is cells[0])
    doc.add_page_break()

    number = 1
    for section_title, questions in SECTIONS:
        doc.add_heading(section_title, level=1)
        intro = doc.add_paragraph()
        intro.paragraph_format.space_after = Pt(10)
        ir = intro.add_run(
            "本部分共{}题。口头回答可根据答辩时间压缩，但不得删除关键边界。".format(len(questions))
        )
        set_run_font(ir, size=9.5, color=(96, 125, 139))
        for item in questions:
            heading = doc.add_heading(f"Q{number:02d}  {item['question']}", level=2)
            heading.paragraph_format.space_before = Pt(10)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(f"提问概率：{item['probability']}")
            set_run_font(r, size=9.5, bold=True, color=(211, 84, 0))
            add_box(doc, "口头回答", item["oral"], "F4F7F9", (28, 84, 122))
            add_box(doc, "小白解释", item["simple"], "EAF6EC", (46, 125, 50))
            add_box(doc, "答辩边界", item["boundary"], "FDEDEC", (183, 28, 28))
            number += 1
        doc.add_page_break()

    doc.add_heading("附录A  已修复问题总览", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("原问题", "解决办法", "当前状态")
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
        set_cell_shading(table.rows[0].cells[i], "1C547A")
        for para in table.rows[0].cells[i].paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=9.5, bold=True, color=(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for issue, solution, status in FIXES:
        cells = table.add_row().cells
        cells[0].text = issue
        cells[1].text = solution
        cells[2].text = status
        if status == "已修复":
            set_cell_shading(cells[2], "E8F5E9")
        elif status in ("基本修复", "常用场景已修复"):
            set_cell_shading(cells[2], "FFF8E1")
        else:
            set_cell_shading(cells[2], "FCE4EC")
        for cell in cells:
            set_cell_margins(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    set_run_font(run, size=9)
    doc.add_page_break()

    doc.add_heading("附录B  答辩红线", level=1)
    add_box(
        doc,
        "原则",
        "宁可明确承认原型边界，也不要把模拟功能、规划功能或器件标称指标说成已经完成的整机能力。",
        "FDEDEC",
        (183, 28, 28),
    )
    for index, text in enumerate(REDLINES, 1):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(text)
        set_run_font(r, size=10.5)

    doc.add_heading("附录C  现场演示检查表", level=1)
    for text in DEMO_CHECKS:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(text)
        set_run_font(r, size=10.5)

    doc.add_heading("最后一分钟速记", level=1)
    add_box(
        doc,
        "必须背熟",
        "项目介绍；核心创新；星火1号与STM32F407选型；为什么不用ESP32；为什么用RT-Thread；燃气安全优先级；MQ-5为什么不能测ppm；精度如何证明；真实端云链路；断网后怎么办；真实与模拟执行器边界；如何证明不是模拟数据。",
        "E8F3F8",
        (17, 94, 130),
    )

    doc.core_properties.title = "AQMV2复赛答辩问题与参考答案"
    doc.core_properties.subject = "RT-Thread星火1号 STM32F407 嵌入式竞赛答辩题库"
    doc.core_properties.author = "AQMV2 Team"
    doc.core_properties.keywords = "AQMV2, RT-Thread, STM32F407, 星火1号, OneNET, 答辩"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
