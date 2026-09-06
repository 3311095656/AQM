from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from generate_defense_qa_docx import (
    add_box,
    add_page_number,
    add_toc,
    set_cell_margins,
    set_cell_shading,
    set_repeat_table_header,
    set_run_font,
)


OUTPUT = Path(__file__).with_name("AQMV2_项目概述与软硬件架构背诵稿.docx")


PROJECT_30S = (
    "我们做的是一套基于星火1号、STM32F407和RT-Thread的智慧环境监测与安全联动原型。"
    "系统采集室内外温湿度、光照、人员存在和可燃气体阈值状态，能够联动风扇、窗户和声光报警，"
    "并将设备真实数据上传OneNET供网页查看和控制。项目重点解决只监测不处理、多种控制指令相互冲突，"
    "以及断网后安全功能失效三个问题，最终形成采集、判断、执行、上报和反馈的完整闭环。"
)


PROJECT_90S = (
    "我们的项目是一套基于星火1号开发板、STM32F407和RT-Thread的智慧环境监测与安全联动系统。"
    "它不是单纯显示几个传感器数值，而是把环境采集、本地判断、设备执行、云端上报和网页控制连接成完整闭环。\n\n"
    "输入端使用AHT10、DHT22、AP3216C、LD2410C和MQ-5，分别获取室内外温湿度、光照、人员存在和燃气是否超过设定阈值。"
    "输出端真实控制风扇、SG90窗户舵机、蜂鸣器和报警灯，空调、加湿器和除湿器目前使用WS2812灯光模拟。\n\n"
    "我们主要解决三个问题。第一，传统监测设备只能发现异常，不能主动处理，我们加入了本地自动控制。"
    "第二，远程、手动、定时和安全指令可能冲突，因此把燃气安全设为最高优先级；发生报警时强制通风、开窗和声光报警，"
    "并在STM32端拒绝危险的关闭命令。第三，安全功能不能依赖云端，所以核心判断都运行在设备端，即使网络断开仍能本地工作。\n\n"
    "软件采用RT-Thread分层和多任务架构，设备通过Paho MQTT直连OneNET；网页通过Node代理OneNET HTTPS API查询数据和下发命令。"
    "当前已经实现真实采集、自动联动、云端记录和远程控制。需要说明的是，MQ-5目前只做阈值报警，部分大功率执行器采用灯光模拟，"
    "整个系统属于竞赛功能验证样机，不是经过计量、燃气或防爆认证的产品。"
)


PROJECT_3MIN = (
    "本项目名为AQMV2，是一套面向室内环境管理和燃气安全联动的嵌入式原型。主控使用星火1号开发板上的STM32F407ZGT6，"
    "操作系统采用RT-Thread。系统通过AHT10、DHT22、AP3216C、LD2410C和MQ-5采集室内外温湿度、光照、人员存在和燃气阈值状态，"
    "在LCD和OLED本地显示，并控制风扇、SG90窗户舵机、蜂鸣器和报警灯；空调、加湿和除湿功能目前由WS2812灯光模拟。\n\n"
    "系统的第一项价值是从只监测提升为闭环处理。STM32把传感器数据整理成统一结构，在本地根据模式、阈值和人体存在进行判断，"
    "随后控制执行器并上报状态。第二项价值是建立明确的安全优先级。MQ-5由独立高优先级线程以约200毫秒周期检测，报警时声光立即动作，"
    "主循环强制开窗、开风扇并关闭普通环境调节状态；MQTT命令入口还会拒绝关风扇、关窗、关闭蜂鸣器等危险命令。"
    "第三项价值是本地自治。即使Wi-Fi、OneNET或网页暂时不可用，本地采集、显示和燃气安全联动仍可继续运行。\n\n"
    "软件按底层驱动、设备模块、控制逻辑、通信服务和云端网页五层组织。关键并发单元包括MQ-5安全线程、舵机时序线程、网络线程和主控制循环。"
    "主线程与网络线程通过互斥锁交换完整传感器快照。固件使用Paho MQTT直接连接OneNET，网页不保存OneNET Token，而是请求本地Node服务，"
    "由Node携带Token调用OneNET HTTPS API。Node不是MQTT Broker，也不转发固件的MQTT数据。\n\n"
    "从功能验证看，系统已经具备真实采集、本地决策、执行器联动、属性上报、历史查询和远程控制能力；普通环境传感器初始化失败不会拖垮燃气安全线程，"
    "系统会暂停相关自动控制并每30秒重试初始化。当前边界也很明确：没有自制PCB，使用开发板和模块搭建；MQ-5不能给出准确ppm；"
    "LD2410C只提供有人或无人；空调与湿控为灯光模拟；没有离线遥测补传、产品级硬件看门狗和安全认证。"
)


PROBLEMS = [
    (
        "只监测、不处理",
        "传统环境监测通常只显示数值，异常后仍依赖人工处理。",
        "把采集、规则判断和执行器控制放在STM32本地。",
        "形成“感知—判断—执行—反馈”闭环。",
    ),
    (
        "控制来源互相冲突",
        "手动、自动、定时和远程命令可能覆盖安全动作。",
        "建立燃气安全最高优先级，并在MQTT入口拒绝危险命令。",
        "报警期间强制通风、开窗和声光报警，危险命令返回403。",
    ),
    (
        "安全功能依赖网络",
        "如果关键判断只在云端，断网时系统可能失去安全能力。",
        "核心安全判断和执行保留在STM32端。",
        "断网后本地采集、显示、报警和安全联动仍可运行。",
    ),
    (
        "端云链路真假不清",
        "初版正式固件连接OneNET，网页却依赖本地Broker模拟数据。",
        "统一为固件直连OneNET MQTT、Node代理OneNET HTTPS API。",
        "网页数据和控制都围绕同一台OneNET真实设备形成闭环。",
    ),
    (
        "普通故障扩散到安全模块",
        "普通温湿度或光照传感器失败曾可能让主程序提前退出。",
        "MQ-5安全线程优先启动；普通传感器初始化失败后降级并重试。",
        "初始化失败时暂停相关自动控制，安全功能保持运行，每30秒重试初始化。",
    ),
]


EFFECTS = [
    (
        "真实采集与本地显示",
        "多类传感器形成统一数据结构，LCD/OLED显示环境和状态。",
        "改变光照、温湿度或人体状态，观察屏幕和串口同步变化。",
        "滤波不等于计量校准；没有标准仪器对比就不宣称整机精度。",
    ),
    (
        "本地自动控制",
        "根据室内温湿度、光照、人体存在和工作模式控制设备。",
        "改变输入条件，观察窗户角度、风扇或WS2812状态变化。",
        "阈值是原型演示参数，空调和湿控目前为灯光模拟。",
    ),
    (
        "燃气安全联动",
        "报警时蜂鸣器和红灯动作，强制开风扇、开窗并关闭普通调节状态。",
        "使用板载安全模拟按键触发，现场观察四类动作。",
        "MQ-5只提供数字阈值状态，不是ppm测量，也不是认证报警器。",
    ),
    (
        "安全命令仲裁",
        "报警期间拒绝关风扇、关窗、关闭蜂鸣器和开启普通调节设备。",
        "报警时从网页发送危险命令，检查403回复和执行器保持安全状态。",
        "软件回复不能替代带位置、转速或电流传感器的物理反馈。",
    ),
    (
        "真实端云闭环",
        "STM32通过MQTT上报OneNET，Node通过HTTPS API为网页查询和控制。",
        "同时核对实物、串口、OneNET时间戳和网页属性变化。",
        "现场必须关闭旧模拟脚本并完成最终真机复测。",
    ),
    (
        "断网与故障降级",
        "网络异常不影响本地安全；普通传感器初始化失败时暂停相关自动控制并重试。",
        "断开热点观察本地功能，再恢复网络确认重新上线；启动时断开普通传感器可验证初始化降级。",
        "当前没有离线遥测队列；30秒重试不覆盖所有运行期传感器故障。",
    ),
]


SOFTWARE_LAYERS = [
    (
        "1. BSP与驱动层",
        "STM32 HAL、RT-Thread BSP、GPIO、I2C、SPI、LCD、WLAN设备框架",
        "向上提供统一硬件访问接口。",
    ),
    (
        "2. 设备模块层",
        "aht10_app、dht22_app、ap3216c_app、mq5_app、radar_app、fan_app、servo_app、lcd_app、oled_app、ws2812b_app",
        "把每类传感器和执行器封装成独立模块。",
    ),
    (
        "3. 数据与控制层",
        "sensor_app、自动控制、手动模式、schedule_app、燃气安全仲裁",
        "完成数据汇总、模式判断、冲突处理和执行决策。",
    ),
    (
        "4. 通信服务层",
        "wifi_app、RW007、Paho MQTT、mqtt_app、NTP、OneNET物模型",
        "负责联网、时间同步、属性上报、命令接收和回复。",
    ),
    (
        "5. 云端与交互层",
        "OneNET、Node/Express、浏览器网页",
        "负责云端存储、API代理、历史查询和远程操作。",
    ),
]


THREADS = [
    (
        "MQ-5安全线程",
        "约200ms / 优先级8",
        "读取MQ-5数字输出和安全模拟按键，控制蜂鸣器与红灯。",
        "最先初始化，独立于普通环境传感器。",
    ),
    (
        "舵机线程",
        "约20ms / 优先级10",
        "生成SG90软件PWM并维护窗户角度。",
        "当前为软件PWM，工程化可改硬件定时器PWM。",
    ),
    (
        "网络线程",
        "1s循环 / 优先级12",
        "处理本地按键、Wi-Fi/MQTT状态、NTP和周期属性上报。",
        "联网后大约每2秒组织一次属性上报。",
    ),
    (
        "主控制循环",
        "约1s循环",
        "传感器采集、双屏显示、计划检查、自动控制和安全强制动作。",
        "与网络线程通过互斥锁交换传感器快照。",
    ),
    (
        "Paho/WLAN内部任务",
        "由软件包管理",
        "完成MQTT收发、心跳、断线重连和网络事件处理。",
        "业务代码不自行实现完整MQTT协议栈。",
    ),
]


HARDWARE = [
    ("AHT10", "室外温湿度", "I2C3", "真实采集"),
    ("DHT22", "室内温湿度", "GPIO单总线 / PA8", "真实采集"),
    ("AP3216C", "环境光照", "I2C2", "当前只用光照"),
    ("LD2410C", "人体存在", "数字OUT / PE14", "只有有人/无人"),
    ("MQ-5", "可燃气体阈值", "数字DO / PA0", "无ADC、无ppm"),
    ("LCD", "室外数据及网络状态", "板载BSP驱动", "真实显示"),
    ("OLED", "室内、燃气和人体状态", "I2C4", "真实显示"),
    ("L9110S + 风扇", "通风", "PG2 / PG3", "真实执行器"),
    ("SG90舵机", "模拟窗户开度", "软件PWM / PE13", "真实动作、无位置反馈"),
    ("蜂鸣器 + 红灯", "声光报警", "PB0 / PF12", "真实执行器"),
    ("WS2812", "模拟空调、加湿、除湿", "单线时序 / PG6", "灯光模拟"),
    ("RW007", "Wi-Fi联网", "SPI2", "连接路由器和OneNET"),
    ("8MB NOR Flash", "预留存储资源", "板载资源", "当前未使用"),
]


QA = [
    (
        "为什么称为硬件架构，没有自制PCB也算硬件吗？",
        "开发板、传感器、执行器、驱动模块、供电和连接线构成了真实硬件系统，所以可以称为硬件架构或硬件集成样机。"
        "但我们没有自制PCB，不能称为自主设计控制板。当前硬件工作主要是器件选型、接口分配、供电与驱动适配和整机联调。",
    ),
    (
        "解决效果能否量化？",
        "目前能够量化的是代码调度参数，例如MQ-5线程约200毫秒检查一次、主循环约1秒、联网后约2秒组织一次上报。"
        "这些是软件周期，不等于经过仪器测得的端到端时延。传感器精度、长期稳定性和误报率仍需要标准仪器、重复试验和长期运行数据证明。",
    ),
    (
        "为什么使用RT-Thread？",
        "系统同时存在安全检测、舵机时序、网络通信、显示和普通控制等不同周期任务。RT-Thread让关键任务可以独立调度，"
        "并通过优先级、互斥锁和定时器降低模块耦合。裸机也能实现，但状态机会更复杂，网络任务更容易影响安全任务。",
    ),
    (
        "Node、Paho、MQTT和Broker分别是什么？",
        "Paho是STM32使用的MQTT客户端库；MQTT是通信协议；Broker是按主题转发MQTT消息的服务器；OneNET提供云端Broker和物模型；"
        "Node只保管网页侧API Token并代理OneNET HTTPS API，不是Broker，也不参与固件MQTT数据转发。",
    ),
    (
        "断网后怎么办？",
        "安全判断和执行在STM32本地，所以断网后仍能采集、显示和执行燃气安全联动。失去的是网页远程控制、云端查询和断网期间的历史上报；"
        "网络恢复后Wi-Fi与MQTT会尝试重连，但当前没有离线数据补传。",
    ),
    (
        "怎样证明不是模拟数据？",
        "现场改变真实光照和人体状态，同时核对实物、串口、OneNET时间戳和网页变化；再从网页控制真实风扇和舵机。"
        "燃气演示使用安全模拟按键，不使用危险气体；正式演示不运行任何本地模拟发布脚本。",
    ),
]


RED_LINES = [
    "不要说MQ-5能测准确ppm；当前只能判断是否超过模块阈值。",
    "不要说LD2410C能测距离、人数或轨迹；当前只读取有人/无人。",
    "不要说已经控制真实空调、加湿器和除湿器；当前由WS2812模拟状态。",
    "不要说自主设计了PCB；准确说法是基于开发板和模块的硬件系统集成。",
    "不要把Paho说成Broker，也不要把Node说成MQTT中转站。",
    "不要把QoS1说成执行器一定成功；它只说明Broker确认收到消息，且可能重复。",
    "不要把滤波说成标定，也不要把器件手册精度当成整机实测精度。",
    "不要说断网数据能够补传；当前只保证本地功能继续运行并在网络恢复后继续上报。",
    "不要说所有运行期传感器故障都能自动恢复；当前30秒重试主要覆盖初始化失败。",
    "不要把定时计划说成已完成端到端验证；当前网页时段enabled字段仍需修正并复测。",
    "不要说Node已经实现完整安全认证；Token仍是服务端静态配置，代理也没有登录与访问控制。",
    "不要说全链路TLS；Node访问OneNET使用HTTPS，但固件MQTT 1883和本地网页HTTP不是全链路加密。",
    "不要说当前软件看门狗能检查每个线程；工程化仍需硬件IWDG加各线程心跳。",
    "不要声称产品已经通过燃气、计量、防爆或长期可靠性认证。",
]


DEMO_CHECKLIST = [
    "烧录最新固件并冷启动，确认LCD、OLED和串口日志正常。",
    "按KEY1连接Wi-Fi和OneNET，确认设备状态变为在线。",
    "改变光照和人体存在状态，核对实物、串口、OneNET和网页同步变化。",
    "从网页控制真实风扇、舵机窗户和蜂鸣器，观察动作及属性更新。",
    "用板载安全模拟按键触发燃气报警，验证蜂鸣器、红灯、开窗和风扇。",
    "报警期间从网页尝试关闭风扇或窗户，确认固件拒绝危险命令。",
    "断开热点，确认本地显示和安全联动继续运行；恢复后确认重新上线。",
    "关闭旧版test.js等模拟消息脚本，准备串口、OneNET和网页截图及备用录像。",
]


def add_bullets(doc, items, color=(38, 50, 56)):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.5, color=color)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(item)
        set_run_font(r, size=10.5)


def add_mono_block(doc, text, fill="EAF2F8"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=180, bottom=130, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=9.5, color=(29, 53, 65))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, header_fill="176B87"):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for idx, value in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = value
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, size=9.2, bold=True, color=(255, 255, 255))
    set_repeat_table_header(table.rows[0])

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = str(value)
            if row_index % 2:
                set_cell_shading(cells[idx], "F5F8FA")
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[idx], top=90, start=100, bottom=90, end=100)
            for para in cells[idx].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.1
                for run in para.runs:
                    set_run_font(run, size=8.8)
        if widths:
            for idx, width in enumerate(widths):
                cells[idx].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in (
        ("Title", 28, (18, 73, 95)),
        ("Heading 1", 18, (18, 73, 95)),
        ("Heading 2", 13, (22, 113, 135)),
        ("Heading 3", 11, (62, 78, 86)),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("AQMV2  |  项目概述与软硬件架构背诵稿")
    set_run_font(run, size=8.5, color=(105, 120, 128))
    add_page_number(section.footer.paragraphs[0])

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(52)
    r = p.add_run("AQMV2")
    set_run_font(r, size=36, bold=True, color=(18, 73, 95))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项目概述与软硬件架构背诵稿")
    set_run_font(r, size=23, bold=True, color=(22, 113, 135))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("做了什么 · 解决了什么 · 效果如何 · 代码架构 · 硬件架构")
    set_run_font(r, size=11.5, color=(87, 105, 114))

    cover = doc.add_table(rows=5, cols=2)
    cover.style = "Table Grid"
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = [
        ("开发平台", "星火1号 / STM32F407ZGT6"),
        ("操作系统", "RT-Thread"),
        ("云端架构", "固件直连OneNET MQTT / Node代理OneNET HTTPS API"),
        ("样机性质", "开发板与成熟模块搭建的功能验证样机，无自制PCB"),
        ("整理日期", "2026-07-21"),
    ]
    for row, (left, right) in zip(cover.rows, values):
        row.cells[0].text = left
        row.cells[1].text = right
        set_cell_shading(row.cells[0], "D9EEF2")
        for cell in row.cells:
            set_cell_margins(cell, top=110, start=150, bottom=110, end=150)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                for rr in para.runs:
                    set_run_font(rr, size=10, bold=cell is row.cells[0])

    doc.add_paragraph()
    add_box(
        doc,
        "一句话定位",
        "这不是一台只显示数据的温湿度计，而是一套把真实采集、本地决策、安全联动、云端记录和网页控制连成闭环的竞赛原型。",
        "E8F5E9",
        (38, 112, 72),
    )
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    doc.add_heading("使用方法", level=1)
    add_box(
        doc,
        "最短背诵主线",
        "采集 → 本地判断 → 安全仲裁 → 执行器 → MQTT上云 → Node代理 → 网页反馈",
        "EAF2F8",
        (18, 73, 95),
    )
    add_numbered(
        doc,
        [
            "先背30秒总述，确保任何时候都能完整回答。",
            "再背“做了什么、解决了什么、效果如何”三段，每段只记关键词。",
            "最后理解代码和硬件架构，不必逐字背引脚，但必须说清模块职责。",
            "遇到精度、ppm、PCB、真实家电和安全认证问题时，主动说明原型边界。",
        ],
    )
    doc.add_heading("目录", level=2)
    add_toc(doc)
    doc.add_page_break()

    doc.add_heading("第一部分  三档项目口述", level=1)
    doc.add_heading("1. 30秒总述：必须逐句熟练", level=2)
    add_box(doc, "直接口述", PROJECT_30S, "E8F5E9", (38, 112, 72))
    add_box(
        doc,
        "关键词",
        "星火1号 / STM32F407 / RT-Thread / 多传感器 / 安全联动 / OneNET / 本地自治 / 完整闭环",
        "FFF8E1",
        (148, 103, 0),
    )

    doc.add_heading("2. 90秒完整介绍：答辩开场推荐", level=2)
    add_box(doc, "直接口述", PROJECT_90S, "EAF2F8", (18, 73, 95))

    doc.add_heading("3. 3分钟展开版：裁判要求详细介绍时使用", level=2)
    add_box(doc, "直接口述", PROJECT_3MIN, "F3F6F7", (62, 78, 86))
    doc.add_page_break()

    doc.add_heading("第二部分  项目做了什么", level=1)
    add_box(
        doc,
        "可背回答",
        "本项目完成了五个环节：第一是采集，通过多个传感器获取温湿度、光照、人员存在和燃气阈值状态；"
        "第二是判断，由STM32根据工作模式、环境阈值和安全规则进行决策；第三是执行，控制风扇、窗户舵机、蜂鸣器、报警灯以及模拟环境调节设备；"
        "第四是联网，通过MQTT把设备真实数据上传OneNET；第五是交互，网页可以查看当前数据、历史数据并远程控制设备。"
        "因此它不是单一监测仪，而是一套感知、决策、执行、联网和反馈的嵌入式原型。",
        "E8F5E9",
        (38, 112, 72),
    )
    add_mono_block(
        doc,
        "真实传感器 → STM32采集 → 本地决策 → 安全仲裁 → 真实/模拟执行器\n"
        "                         ↓                         ↑\n"
        "                  OneNET MQTT ←→ OneNET平台 ←HTTPS→ Node ←HTTP→ 网页",
    )
    add_table(
        doc,
        ("环节", "项目实现", "一句话记忆"),
        [
            ("感知", "室内外温湿度、光照、人体存在、燃气阈值", "看见环境"),
            ("决策", "自动、手动、远程和安全优先级；固件另含待复测的定时模块", "判断该做什么"),
            ("执行", "风扇、舵机、蜂鸣器、LED及WS2812模拟设备", "真正做出动作"),
            ("联网", "Paho MQTT直连OneNET，QoS1属性上报和回复", "把状态送上云"),
            ("交互", "Node代理API，网页实时属性、历史和远程控制", "让用户看和管"),
        ],
        widths=(2.2, 10.2, 3.2),
    )

    doc.add_heading("小白理解", level=2)
    add_box(
        doc,
        "房间管家类比",
        "传感器是眼睛和鼻子，STM32是大脑，RT-Thread是任务调度员，风扇和舵机是手脚，MQTT是设备与云端通信的规则，"
        "OneNET是云端数据和消息中心，Node是保管API钥匙并替网页办事的管家，网页是用户控制面板。燃气安全规则是最高级紧急命令。",
        "FFF8E1",
        (148, 103, 0),
    )
    doc.add_page_break()

    doc.add_heading("第三部分  项目解决了什么", level=1)
    add_box(
        doc,
        "可背回答",
        "项目重点解决四类问题：把只监测不处理变为本地自动处理；用安全优先级解决手动、自动、定时和远程命令冲突；"
        "把核心安全逻辑留在STM32端，避免断网后失去安全能力；同时修复初版真机与网页链路分离的问题，让网页围绕OneNET真实设备工作。"
        "此外，普通环境传感器初始化失败后系统会降级而不是拖垮燃气安全功能。",
        "E8F5E9",
        (38, 112, 72),
    )
    add_box(
        doc,
        "四组关键词",
        "只看不管 → 主动处理　|　指令冲突 → 安全优先　|　依赖云端 → 本地自治　|　模拟链路 → 真实闭环",
        "FFF8E1",
        (148, 103, 0),
    )
    add_table(
        doc,
        ("要解决的问题", "原来的风险", "现在的办法", "达到的效果"),
        PROBLEMS,
        widths=(3.1, 4.5, 5.0, 4.0),
    )

    doc.add_heading("最重要的一项修复：真实端云链路", level=2)
    add_mono_block(
        doc,
        "初版：STM32 → OneNET Broker　　网页/Node → 本地Aedes Broker\n"
        "      两条链路没有桥接，网页可显示模拟消息，但不能证明真机闭环。\n\n"
        "现在：STM32 ←MQTT→ OneNET ←HTTPS API→ Node ←HTTP→ 网页\n"
        "      OneNET成为唯一云端数据与控制中心。",
        fill="FDEDEC",
    )
    add_box(
        doc,
        "裁判追问时的结论",
        "初版的问题不是网页完全不能动，而是本地测试脚本能够制造出“看起来有数据、按钮有响应”的局部演示效果，"
        "却没有证明网页命令经过OneNET到达真实STM32。复赛版统一链路后，正式页面不再依赖本地模拟Broker。",
        "FDEDEC",
        (170, 45, 45),
    )
    doc.add_page_break()

    doc.add_heading("第四部分  解决效果怎么样", level=1)
    add_box(
        doc,
        "可背回答",
        "从功能验证角度看，传感器数据能够在本地显示并上传OneNET，网页可以查询设备实际属性和历史；网页控制请求可以经过Node和OneNET到达STM32；"
        "燃气阈值触发后，系统能够打开风扇和窗户并启动声光报警；报警期间，STM32会拒绝危险命令；网络中断不会停止本地安全控制，"
        "普通环境传感器初始化失败时系统也会降级并定期重试。因此项目已经实现功能闭环和安全优先策略，但还不能把功能验证说成计量精度、长期可靠性或产品安全认证。",
        "E8F5E9",
        (38, 112, 72),
    )
    add_table(
        doc,
        ("验证目标", "当前效果", "现场证据", "必须说明的边界"),
        EFFECTS,
        widths=(3.0, 5.0, 4.8, 4.3),
    )

    doc.add_heading("效果应该怎样证明", level=2)
    add_mono_block(
        doc,
        "传感器变化 → 本地屏幕、串口、OneNET、网页一起变化\n"
        "网页点击　 → OneNET下发 → STM32判断 → 真实风扇/舵机动作\n"
        "触发报警　 → 蜂鸣器 + 红灯 + 开窗 + 风扇\n"
        "报警时关风扇 → STM32拒绝命令，保持安全状态\n"
        "断开网络　 → 本地采集和安全联动继续工作",
    )
    add_box(
        doc,
        "关于“实时性”的准确说法",
        "代码中MQ-5线程约每200毫秒检查一次，主控制循环约1秒，联网后约每2秒组织一次属性上报，网页也约每2秒查询一次。"
        "这些是软件调度周期，不等于经过仪器测得的端到端响应时间；网络延迟、传感器响应和舵机机械时间还需要单独测量。",
        "FFF8E1",
        (148, 103, 0),
    )
    add_box(
        doc,
        "效果等级结论",
        "当前达到的是“竞赛原型的功能闭环验证”，还没有达到“产品级计量、长期可靠性、信息安全和燃气防爆认证”。",
        "FDEDEC",
        (170, 45, 45),
    )
    doc.add_page_break()

    doc.add_heading("第五部分  代码架构", level=1)
    add_box(
        doc,
        "可背回答",
        "软件采用分层、模块化和RT-Thread多任务架构。最底层是STM32 BSP、HAL以及GPIO、I2C、SPI和WLAN驱动；上面是各传感器、执行器和显示模块；"
        "再上面是统一数据采集、自动控制、模式管理、燃气安全仲裁以及仍需前端字段修正复测的定时模块；通信层包括RW007、Paho MQTT、OneNET物模型和NTP；"
        "最上层是OneNET、Node API代理和网页。只有安全检测、舵机时序和网络等关键功能使用独立任务，普通采集和控制由主循环统一调度。",
        "E8F5E9",
        (38, 112, 72),
    )
    add_mono_block(
        doc,
        "浏览器网页\n"
        "    ↓ HTTP\n"
        "Node / Express API代理 ──HTTPS──→ OneNET API与历史存储\n"
        "                                      ↕ 平台内部\n"
        "设备应用逻辑 ─→ mqtt_app ─→ Paho MQTT ─→ OneNET MQTT Broker\n"
        "    ↑\n"
        "RT-Thread设备模块与STM32 BSP/HAL",
    )
    add_table(
        doc,
        ("软件层", "主要模块", "职责"),
        SOFTWARE_LAYERS,
        widths=(3.3, 8.3, 5.1),
    )

    doc.add_heading("RT-Thread运行单元", level=2)
    add_table(
        doc,
        ("线程/循环", "周期与优先级", "主要职责", "设计说明"),
        THREADS,
        widths=(3.0, 3.2, 5.3, 5.2),
    )
    add_box(
        doc,
        "为什么不让每个模块都创建线程",
        "传感器和执行器首先按模块封装，但只有需要独立周期、时序或故障隔离的功能才创建线程。这样既保持职责清楚，也避免线程过多造成栈内存浪费和同步复杂度上升。",
        "FFF8E1",
        (148, 103, 0),
    )

    doc.add_heading("核心数据流与命令流", level=2)
    add_mono_block(
        doc,
        "数据流：传感器 → sensor_data_t → 互斥锁缓存 → 显示/控制 → MQTT上报 → OneNET → Node → 网页\n\n"
        "命令流：网页 → Node加入API Token → OneNET HTTPS API → OneNET MQTT → mqtt_app校验\n"
        "        → 燃气安全仲裁 → 执行器 → set_reply与属性上报",
    )
    add_box(
        doc,
        "安全优先级",
        "燃气安全 > 用户控制和固件定时逻辑 > 普通自动舒适性控制。报警期间，安全动作不仅由主循环持续强制，MQTT命令入口也会直接拒绝危险操作；定时功能的网页字段仍需修正复测。",
        "FDEDEC",
        (170, 45, 45),
    )

    doc.add_heading("Node、Paho、MQTT和Broker不要混淆", level=2)
    add_table(
        doc,
        ("名称", "本质", "在本项目中的作用"),
        [
            ("API Token", "身份凭证", "由Node服务端保存，用于调用OneNET HTTPS API"),
            ("Node", "网页服务器与API代理", "托管页面、加入Token、过滤字段、代理查询和控制"),
            ("Paho MQTT", "STM32端客户端库", "实现连接、发布、订阅、心跳和重连"),
            ("MQTT", "通信协议", "规定设备与Broker之间怎样收发消息"),
            ("Broker", "MQTT消息服务器", "按主题接收并转发消息"),
            ("OneNET", "物联网云平台", "提供Broker、物模型、存储和HTTPS API"),
        ],
        widths=(3.0, 4.8, 8.9),
    )
    doc.add_page_break()

    doc.add_heading("第六部分  硬件架构", level=1)
    add_box(
        doc,
        "可背回答",
        "硬件以星火1号开发板上的STM32F407ZGT6为核心。输入端连接室内外温湿度、光照、人体存在和燃气阈值传感器；"
        "显示端使用LCD、OLED和WS2812；执行端包括L9110S风扇驱动、SG90舵机、蜂鸣器和LED；通信端通过SPI2连接RW007 Wi-Fi模块。"
        "当前没有自制PCB，而是用开发板、成熟模块和杜邦线搭建功能验证样机，因此准确说法是硬件系统集成，不能说成自主设计控制板。",
        "E8F5E9",
        (38, 112, 72),
    )
    add_mono_block(
        doc,
        "AHT10 / DHT22 / AP3216C / LD2410C / MQ-5 / 按键\n"
        "                         ↓\n"
        "              星火1号 STM32F407ZGT6\n"
        "       ┌─────────────────┼─────────────────┐\n"
        "       ↓                 ↓                 ↓\n"
        " LCD/OLED/WS2812   风扇/舵机/声光报警   RW007(SPI2)\n"
        "                                             ↓\n"
        "                                      Wi-Fi → OneNET",
    )
    add_table(
        doc,
        ("硬件", "用途", "接口/引脚", "当前边界"),
        HARDWARE,
        widths=(3.0, 4.2, 4.5, 5.0),
    )

    doc.add_heading("真实执行器与模拟执行器边界", level=2)
    add_table(
        doc,
        ("类别", "内容", "答辩表述"),
        [
            ("真实采集", "AHT10、DHT22、AP3216C、MQ-5数字输出、LD2410C OUT", "输入来自真实硬件，但能力受接口选择限制"),
            ("真实动作", "风扇、SG90舵机、蜂鸣器、LED、LCD、OLED", "可以现场观察动作或显示"),
            ("状态模拟", "空调、加湿器、除湿器", "WS2812只验证控制逻辑和接口，不代表接入真实大功率设备"),
            ("无闭环反馈", "风扇转速、舵机位置、负载电流", "当前上报主要是软件命令状态，不是独立物理反馈"),
        ],
        widths=(3.0, 6.0, 7.7),
    )

    doc.add_heading("没有PCB，为什么仍可称为硬件架构", level=2)
    add_box(
        doc,
        "可背回答",
        "开发板、传感器、执行器、驱动模块、供电和连接线共同构成真实硬件系统，因此可以称为硬件架构或硬件集成样机。"
        "但没有自制PCB就不能称为自主硬件控制板。杜邦线方案适合功能验证，工程化阶段还需要专用PCB、稳定连接器、电源保护、滤波、接口防护和EMC验证。",
        "FFF8E1",
        (148, 103, 0),
    )

    doc.add_heading("板载8MB NOR Flash为什么没用", level=2)
    add_box(
        doc,
        "可背回答",
        "当前复赛阶段优先完成安全控制和真实端云闭环，因此8MB NOR Flash尚未进入正式数据流。它不是系统当前必须依赖的资源。"
        "后续可以配合FAL和EasyFlash保存定时配置、故障日志及断网期间的数据队列，实现掉电保持和恢复后补传。",
        "EAF2F8",
        (18, 73, 95),
    )
    doc.add_page_break()

    doc.add_heading("第七部分  高概率追问速答", level=1)
    for index, (question, answer) in enumerate(QA, start=1):
        doc.add_heading(f"{index}. {question}", level=2)
        add_box(doc, "直接回答", answer, "EAF2F8", (18, 73, 95))

    doc.add_heading("第八部分  不能说错的边界", level=1)
    add_box(
        doc,
        "总原则",
        "始终分清三件事：已经实现的能力、竞赛原型的限制、下一步计划。不能用未来计划回答当前已经实现。",
        "FDEDEC",
        (170, 45, 45),
    )
    add_bullets(doc, RED_LINES, color=(145, 35, 35))

    doc.add_heading("第九部分  上场前真机检查", level=1)
    add_numbered(doc, DEMO_CHECKLIST)
    add_box(
        doc,
        "最终收尾句",
        "这个项目的核心价值不在于传感器数量，而在于把真实采集、本地自治、安全优先、执行反馈和端云闭环组合在一起；"
        "当前完成的是可验证的竞赛原型，下一步才是专用PCB、离线存储、硬件看门狗和产品级可靠性设计。",
        "E8F5E9",
        (38, 112, 72),
    )

    doc.core_properties.title = "AQMV2项目概述与软硬件架构背诵稿"
    doc.core_properties.subject = "嵌入式竞赛复赛答辩背诵材料"
    doc.core_properties.author = "AQMV2 Team"
    doc.core_properties.keywords = "AQMV2, STM32F407, RT-Thread, OneNET, 代码架构, 硬件架构, 答辩"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
